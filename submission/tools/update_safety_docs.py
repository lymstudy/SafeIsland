#!/usr/bin/env python3
"""Patch safety docx files with post-TMR inventory statistics."""

from __future__ import annotations

import re
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]


def run_gen_fault_lists():
    script = ROOT / "tools" / "gen_fault_lists.py"
    subprocess.check_call([sys.executable, str(script)], cwd=ROOT)


def parse_summary(path: Path):
    text = path.read_text(encoding="utf-8")
    total = 0
    m = re.search(r"total_rows=(\d+)", text)
    if m:
        total = int(m.group(1))
    by_module = {}
    in_mod = False
    for line in text.splitlines():
        if line.startswith("By module:"):
            in_mod = True
            continue
        if in_mod:
            if line.startswith("  ") and ":" in line:
                mod, cnt = line.strip().split(":", 1)
                by_module[mod.strip()] = int(cnt.strip())
            elif line.strip() == "" or line.startswith("By "):
                in_mod = False
    return total, by_module


def replace_paragraph_containing(doc, needle, new_text):
    for para in doc.paragraphs:
        if needle in para.text:
            para.text = new_text
            return True
    return False


def update_doc4(reg_total, logic_total, by_mod):
    from docx import Document

    path = ROOT / "4-安全机制分析及设计.docx"
    doc = Document(path)

    intro = (
        "本文档从硬件设计角度描述 AXI Safety Island 的安全机制。"
        "2026-07 升级后，关键配置表、控制状态、slot 元数据、故障事件 latch 与顶层 fault 输出均采用 "
        "TMR 三副本 + majority voted 功能输出 + 下一拍 feedback repair；"
        "宽字段辅以 parity signature，反码 shadow 降级为 latent 检测辅助；"
        "关键 voter 使用 tmr_voter_protected 自保护表决。"
        "单点 flip/stuck-at 优先归类为 corrected，mismatch/latent 上报，双点/保护链异常 10 cycle 内 detected。"
    )
    replace_paragraph_containing(doc, "AXI Safety Island 的安全机制", intro)

    # Table 17: register coverage by module (index may vary — search by header)
    for table in doc.tables:
        if not table.rows:
            continue
        header = [c.text.strip() for c in table.rows[0].cells]
        if "寄存器" in "".join(header) or "Memory" in "".join(header):
            if len(table.rows) >= 2 and "config_slave" in table.rows[1].cells[0].text:
                # rebuild rows for top modules
                while len(table.rows) > 1:
                    table._tbl.remove(table.rows[1]._tr)
                order = sorted(by_mod.items(), key=lambda x: -x[1])[:7]
                for mod, cnt in order:
                    row = table.add_row().cells
                    row[0].text = mod
                    row[1].text = f"~{cnt} bit-rows"
                    row[2].text = "TMR+repair/scrub/parity/shadow"
                    row[3].text = "corrected / latent / detected"
                    if len(row) > 4:
                        row[4].text = "post-TMR inventory"
                tot_row = table.add_row().cells
                tot_row[0].text = "合计 (Register_fault_list)"
                tot_row[1].text = str(reg_total)
                break

    concl = (
        f"升级后 Register 故障清单 {reg_total} 行、Logic 故障清单 {logic_total} 行。"
        "TMR+repair 覆盖 config_slave 配置表与控制位、read_engine slot 元数据、core FSM/索引/pending、"
        "fault_detector 事件与 status、top sticky fault 输出。"
        "注错 campaign 统计待故障注入 TB 同步升级后重新跑批；工具链见 tools/gen_fault_lists.py、analyze_fi_report.py。"
    )
    if doc.paragraphs:
        doc.paragraphs[-1].text = concl

    doc.save(path)
    print(f"Updated {path}")


def update_doc5(reg_total, logic_total):
    from docx import Document

    path = ROOT / "5-注错仿真计划.docx"
    doc = Document(path)

    replace_paragraph_containing(
        doc,
        "每 case 归类为",
        "每 case 归类为: 已纠正 (corrected)、已探知 (detected)、潜伏 (latent)、未探知 (undetected)。"
        "corrected=功能输出正确(允许 latent=1)；detected=10 cycle 内 fault_detect 或 safety_island_fault_detect；"
        "latent=功能正确但 mismatch/latent 标志置位；undovered=输出错误且未上报。",
    )
    replace_paragraph_containing(
        doc,
        "保护概率",
        "保护概率 = (已纠正 + 已探知 + 潜伏) / 错误总数；注错分母使用 fault_campaign/Register_fault_list.csv + Logic_fault_list.csv。",
    )
    replace_paragraph_containing(
        doc,
        "fault_campaign/Register_fault_list.csv",
        f"fault_campaign/Register_fault_list.csv — {reg_total} 行 (post-TMR 寄存器清单)\n"
        f"fault_campaign/Logic_fault_list.csv — {logic_total} 行 (post-TMR 逻辑清单, bit 级)\n"
        "fault_campaign/legacy_logic_signal_list.csv — 459 信号级逻辑源清单 (生成输入)\n"
        "fault_campaign/fault_smoke_tmr.csv — 14 项 smoke (升级后首批验证)",
    )
    replace_paragraph_containing(
        doc,
        "6. 验证结果摘要",
        "6. 验证结果摘要 (待重新跑批)\n"
        "RTL 安全机制已升级；注错 TB 与 campaign 同步更新后执行: "
        "cd submission/scripts && bash run_fault.sh && python ../tools/analyze_fi_report.py "
        "--input ../sim/fault_injection/reports/fault_injection_report.csv",
    )

    for table in doc.tables:
        if not table.rows:
            continue
        hdr = [c.text for c in table.rows[0].cells]
        if "结果" in "".join(hdr) or "基线" in "".join(hdr):
            if len(table.rows) >= 2:
                table.rows[-1].cells[0].text = "Post-TMR campaign"
                if len(table.rows[-1].cells) > 1:
                    table.rows[-1].cells[1].text = "TBD (pending FI upgrade)"
                if len(table.rows[-1].cells) > 2:
                    table.rows[-1].cells[2].text = "TBD"

    doc.save(path)
    print(f"Updated {path}")


def main():
    run_gen_fault_lists()
    reg_total, by_mod = parse_summary(ROOT / "fault_campaign" / "Register_fault_summary.txt")
    _, logic_by = parse_summary(ROOT / "fault_campaign" / "Logic_fault_summary.txt")
    logic_total = sum(logic_by.values()) if logic_by else 0
    if not logic_total:
        logic_total = len(
            (ROOT / "fault_campaign" / "Logic_fault_list.csv").read_text(encoding="utf-8").splitlines()
        ) - 1

    update_doc4(reg_total, logic_total, by_mod)
    update_doc5(reg_total, logic_total)

    # refresh metrics CSV without FI (pending)
    subprocess.check_call([sys.executable, str(ROOT / "tools" / "gen_safety_report.py")], cwd=ROOT)


if __name__ == "__main__":
    main()
